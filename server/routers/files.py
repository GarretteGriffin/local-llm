from __future__ import annotations

from io import BytesIO
from typing import Literal, Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

router = APIRouter(prefix="/files", tags=["files"])


class CreateOfficeFileRequest(BaseModel):
    kind: Literal["word", "excel", "powerpoint"]
    filename: Optional[str] = Field(default=None, description="Optional base filename (without extension)")


def _sanitize_filename(base: str) -> str:
    base = (base or "").strip()
    if not base:
        return "New File"
    # Avoid path traversal and illegal filename characters.
    for ch in ("/", "\\", ":", "*", "?", '"', "<", ">", "|"):
        base = base.replace(ch, " ")
    base = " ".join(base.split())
    return base[:80] if base else "New File"


@router.post("/create")
async def create_office_file(req: CreateOfficeFileRequest) -> Response:
    """Create a blank Office file and return it as a download."""

    kind = req.kind
    base = _sanitize_filename(req.filename or "New File")

    if kind == "word":
        try:
            from docx import Document
        except Exception as e:
            raise HTTPException(status_code=500, detail="python-docx is not installed") from e

        doc = Document()
        bio = BytesIO()
        doc.save(bio)
        data = bio.getvalue()
        filename = f"{base}.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    elif kind == "excel":
        try:
            from openpyxl import Workbook
        except Exception as e:
            raise HTTPException(status_code=500, detail="openpyxl is not installed") from e

        wb = Workbook()
        bio = BytesIO()
        wb.save(bio)
        data = bio.getvalue()
        filename = f"{base}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    elif kind == "powerpoint":
        try:
            from pptx import Presentation
        except Exception as e:
            raise HTTPException(status_code=500, detail="python-pptx is not installed") from e

        prs = Presentation()
        bio = BytesIO()
        prs.save(bio)
        data = bio.getvalue()
        filename = f"{base}.pptx"
        media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    else:
        raise HTTPException(status_code=400, detail="Unsupported file kind")

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
    }
    return Response(content=data, media_type=media_type, headers=headers)
